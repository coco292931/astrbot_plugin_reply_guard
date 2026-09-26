"""关键词表：从 xlsx 读取 A 列关键词、B 列回复内容，后行覆盖前行。"""

from __future__ import annotations

import os
import time
from typing import Any

from astrbot.api import logger


def _cell_text(value: Any) -> str:
    """单元格转字符串，空白和 None 一律算空。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


class KeywordStore:
    """xlsx 关键词表，支持定时热重载与精准匹配。"""

    def __init__(
        self,
        path: str,
        *,
        sheet: str = "",
        reload_interval: int = 60,
        skip_header: bool = False,
    ) -> None:
        self.path = (path or "").strip()
        self.sheet = (sheet or "").strip()
        self.reload_interval = max(0, int(reload_interval or 0))
        self.skip_header = bool(skip_header)
        self._mapping: dict[str, str] = {}
        self._mtime: float = 0.0
        self._last_check: float = 0.0
        self._error: str = ""
        self.rows: int = 0
        self.skipped: int = 0

    @property
    def count(self) -> int:
        return len(self._mapping)

    @property
    def error(self) -> str:
        return self._error

    def snapshot(self) -> dict[str, str]:
        return dict(self._mapping)

    def load(self, *, force: bool = False) -> bool:
        """读表。返回是否真的重新加载了。"""
        if not self.path:
            self._error = "未配置关键词表路径"
            return False

        if not os.path.isfile(self.path):
            self._error = f"关键词表不存在: {self.path}"
            return False

        try:
            mtime = os.path.getmtime(self.path)
        except OSError as e:
            self._error = f"读取文件时间失败: {e}"
            return False

        if not force and mtime == self._mtime and self._mapping:
            return False

        try:
            import openpyxl
        except ImportError:
            self._error = "缺少依赖 openpyxl，请先安装"
            return False

        try:
            wb = openpyxl.load_workbook(self.path, data_only=True, read_only=True)
        except Exception as e:
            self._error = f"打开 xlsx 失败: {e}"
            return False

        try:
            if self.sheet and self.sheet in wb.sheetnames:
                ws = wb[self.sheet]
            else:
                ws = wb.worksheets[0]

            mapping: dict[str, str] = {}
            rows = 0
            skipped = 0
            for row in ws.iter_rows(values_only=True):
                rows += 1
                if self.skip_header and rows == 1:
                    continue
                if not row:
                    skipped += 1
                    continue
                keyword = _cell_text(row[0] if len(row) > 0 else None)
                content = _cell_text(row[1] if len(row) > 1 else None)
                if not keyword or not content:
                    skipped += 1
                    continue
                mapping[keyword] = content

            wb.close()
        except Exception as e:
            self._error = f"解析 xlsx 失败: {e}"
            logger.warning(f"[reply_guard] 解析关键词表失败: {e}")
            try:
                wb.close()
            except Exception:
                pass
            return False

        self._mapping = mapping
        self._mtime = mtime
        self.rows = rows
        self.skipped = skipped
        self._error = ""
        logger.info(
            f"[reply_guard] 关键词表已加载: {self.path} "
            f"有效 {len(mapping)} 条 / 共 {rows} 行，跳过 {skipped} 行"
        )
        return True

    def reload_if_needed(self) -> None:
        """按间隔检查文件是否变化，变化就重新加载。"""
        if self.reload_interval <= 0:
            return
        now = time.time()
        if now - self._last_check < self.reload_interval:
            return
        self._last_check = now
        try:
            if os.path.isfile(self.path) and os.path.getmtime(self.path) != self._mtime:
                self.load()
        except OSError:
            return

    def match(self, text: str) -> str | None:
        """精准匹配，整条消息完全等于关键词才算命中。"""
        if not text:
            return None
        key = text.strip()
        if not key:
            return None
        return self._mapping.get(key)
